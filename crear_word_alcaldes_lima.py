#!/usr/bin/env python3
"""Genera un informe Word a partir de la extracción de alcaldes de Lima."""

from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "salida" / "alcaldes_lima_completo.json"
OUTPUT = ROOT / "salida" / "Informe_alcaldes_Lima_ERM_2026.docx"

BRAND = "8B1E2D"
BRAND_DARK = "55111B"
LIGHT = "F5EDEF"
GRAY = "60646C"

LABELS = {
    "id": "ID",
    "dni": "DNI",
    "nombre": "Nombre completo",
    "edad": "Edad",
    "genero": "Género",
    "partido": "Organización política",
    "partidoId": "ID de organización (JNE)",
    "region": "Región",
    "foto": "URL de fotografía",
    "categoria": "Categoría",
    "cargo": "Cargo",
    "estado": "Estado",
    "numero": "Número",
    "proceso": "Proceso electoral",
    "jurisdiccion": "Jurisdicción",
    "slug": "Código del proceso",
    "fecha": "Fecha",
    "titulo": "Título o grado",
    "universidad": "Institución educativa",
    "anio": "Año",
    "fuente": "Fuente",
    "organizacion": "Organización",
    "elegido": "Elegido/a",
    "ambito": "Ámbito",
    "fechaInicio": "Fecha de inicio",
    "fechaFin": "Fecha de fin",
    "assetType": "Tipo de bien",
    "description": "Descripción",
    "vehiclePlate": "Placa",
    "address": "Dirección",
    "value": "Valor declarado",
    "comment": "Comentario",
    "quantity": "Cantidad",
    "cargo": "Cargo",
    "empleador": "Empleador",
    "anioInicio": "Año de inicio",
    "anioFin": "Año de fin",
    "ubicacion": "Ubicación",
    "ruc": "RUC",
    "businessName": "Razón social",
    "commercialName": "Nombre comercial",
    "contributorType": "Tipo de contribuyente",
    "contributorStatus": "Estado del contribuyente",
    "contributorCondition": "Condición del contribuyente",
    "registrationDate": "Fecha de inscripción",
    "hasCoactiveDebt": "Tiene deuda coactiva",
    "coactiveDebts": "Deudas coactivas",
    "amount": "Monto",
    "category": "Categoría",
    "contractingEntity": "Entidad contratante",
    "currency": "Moneda",
    "endDate": "Fecha de término",
    "isConsortium": "Participación en consorcio",
    "oeceCode": "Código OECE",
    "isForeign": "Empresa extranjera",
    "numContracts": "Número de contratos",
    "contractsAmount": "Monto de contratos",
    "numPenalties": "Número de sanciones",
    "sharePercentage": "Porcentaje de participación",
    "relations": "Relaciones declaradas",
    "tcpSanctions": "Sanciones del Tribunal de Contrataciones",
    "coactiveDebtAmount": "Monto de deuda coactiva",
    "licenseCategory": "Categoría de licencia",
    "licenseExpiry": "Vencimiento de licencia",
    "licenseStatus": "Estado de licencia",
    "accumulatedPoints": "Puntos acumulados",
    "seriousViolations": "Infracciones graves",
    "verySeriousViolations": "Infracciones muy graves",
    "numPapeletas": "Número de papeletas",
    "papeletas": "Papeletas",
    "numTramites": "Número de trámites",
    "numBonificaciones": "Bonificaciones",
    "caseNumber": "Expediente",
    "caseType": "Tipo de proceso",
    "crime": "Delito o materia",
    "sentence": "Sentencia",
    "sentenceDate": "Fecha de sentencia",
    "court": "Órgano jurisdiccional",
    "modality": "Modalidad",
    "complianceStatus": "Estado de cumplimiento",
    "notes": "Notas",
    "barAssociation": "Colegio profesional",
    "barMembershipNumber": "Número de colegiatura",
    "fineType": "Tipo de sanción",
    "institution": "Institución",
    "isActive": "Sanción vigente",
    "registrationNumber": "Número de registro",
    "sanctionDate": "Fecha de sanción",
    "sanctionEndDate": "Fin de sanción",
}


def label(key: str) -> str:
    return LABELS.get(key, key.replace("_", " ").strip().capitalize())


def display(value: Any) -> str:
    if value is None or value == "":
        return "No publicado"
    if value is True:
        return "Sí"
    if value is False:
        return "No"
    if isinstance(value, float):
        return f"{value:,.2f}"
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, separators=(", ", ": "))
    return str(value)


def shade(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_text(cell, text: str, *, bold: bool = False, color: str | None = None, size: int = 9) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_key_value_table(document: Document, data: dict[str, Any], keys: list[str] | None = None) -> None:
    selected = keys or list(data.keys())
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for key in selected:
        if key not in data:
            continue
        row = table.add_row().cells
        shade(row[0], LIGHT)
        set_cell_text(row[0], label(key), bold=True, color=BRAND_DARK)
        set_cell_text(row[1], display(data.get(key)))
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_records(document: Document, records: Any, empty_text: str = "Sin información publicada.") -> None:
    if records is None or records == [] or records == {}:
        paragraph = document.add_paragraph(empty_text)
        paragraph.style = document.styles["Intense Quote"]
        return
    if isinstance(records, dict):
        add_key_value_table(document, records)
        return
    if not isinstance(records, list):
        document.add_paragraph(display(records))
        return
    for index, item in enumerate(records, 1):
        if len(records) > 1:
            paragraph = document.add_paragraph()
            run = paragraph.add_run(f"Registro {index}")
            run.bold = True
            run.font.color.rgb = RGBColor.from_string(BRAND)
        if isinstance(item, dict):
            add_key_value_table(document, item)
        else:
            document.add_paragraph(display(item))


def add_heading(document: Document, text: str, level: int = 2) -> None:
    heading = document.add_heading(text, level=level)
    heading.paragraph_format.keep_with_next = True


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Página ")
    run.font.size = Pt(8)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run._r.addnext(field)


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9)
    normal.paragraph_format.space_after = Pt(4)

    for style_name, size, color in [
        ("Title", 28, BRAND_DARK),
        ("Heading 1", 18, BRAND_DARK),
        ("Heading 2", 13, BRAND),
        ("Heading 3", 11, GRAY),
    ]:
        style = styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)

    header = section.header.paragraphs[0]
    header.text = "REVISIÓN DE CANDIDATURAS · LIMA · ERM 2026"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.name = "Aptos"
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor.from_string(GRAY)
    add_page_number(section.footer.paragraphs[0])


def add_cover(document: Document, metadata: dict[str, Any]) -> None:
    document.add_paragraph("\n\n\n")
    title = document.add_heading("Candidatos a la Alcaldía de Lima", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.add_paragraph("Elecciones Regionales y Municipales 2026")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].bold = True
    subtitle.runs[0].font.size = Pt(16)
    subtitle.runs[0].font.color.rgb = RGBColor.from_string(BRAND)

    document.add_paragraph("\n")
    scope = document.add_paragraph(metadata.get("alcance", ""))
    scope.alignment = WD_ALIGN_PARAGRAPH.CENTER
    scope.runs[0].font.size = Pt(12)

    extracted = metadata.get("fecha_extraccion", "")
    try:
        dt = datetime.fromisoformat(extracted)
        extracted = dt.strftime("%d/%m/%Y %H:%M")
    except (TypeError, ValueError):
        pass
    info = document.add_paragraph(
        f"{metadata.get('cantidad', 0)} perfiles · Extracción: {extracted}\n"
        f"Fuente: {metadata.get('fuente', '')}"
    )
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in info.runs:
        run.font.color.rgb = RGBColor.from_string(GRAY)

    note = document.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.space_before = Pt(60)
    note.add_run(
        "Documento de investigación. Los datos reflejan lo publicado por la fuente "
        "al momento de la extracción y pueden variar durante el proceso electoral."
    ).italic = True
    document.add_page_break()


def add_summary(document: Document, candidates: list[dict[str, Any]]) -> None:
    add_heading(document, "Resumen de candidaturas", 1)
    paragraph = document.add_paragraph(
        "El informe incluye todas las candidaturas a alcalde provincial de Lima encontradas en la fuente, "
        "incluyendo estados no definitivos o desfavorables."
    )
    paragraph.paragraph_format.space_after = Pt(10)

    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, text in zip(table.rows[0].cells, ["N.º", "Candidato/a", "Organización política", "Estado"]):
        shade(cell, BRAND)
        set_cell_text(cell, text, bold=True, color="FFFFFF")

    for index, record in enumerate(candidates, 1):
        listed = record["listado"]
        cells = table.add_row().cells
        values = [str(index), listed.get("name", ""), listed.get("party", ""), listed.get("status", "")]
        for cell, value in zip(cells, values):
            set_cell_text(cell, value, size=8)
        if index % 2 == 0:
            for cell in cells:
                shade(cell, "FAF7F8")
    document.add_page_break()


def add_candidate(document: Document, record: dict[str, Any], number: int, total: int) -> None:
    candidate = record.get("candidato") or {}
    detail = record.get("detalle") or {}
    title = document.add_heading(candidate.get("nombre", "Candidato/a"), 1)
    title.paragraph_format.space_after = Pt(2)
    kicker = document.add_paragraph(
        f"Perfil {number} de {total} · {candidate.get('partido', 'Organización no publicada')}"
    )
    kicker.runs[0].bold = True
    kicker.runs[0].font.color.rgb = RGBColor.from_string(BRAND)

    add_heading(document, "Datos generales", 2)
    general = {
        "id": candidate.get("id"),
        "dni": candidate.get("dni"),
        "edad": candidate.get("edad"),
        "genero": candidate.get("genero"),
        "partido": candidate.get("partido"),
        "partidoId": candidate.get("partidoId"),
        "region": candidate.get("region"),
        "foto": candidate.get("foto"),
        "perfil": record.get("url"),
    }
    LABELS["perfil"] = "URL del perfil"
    add_key_value_table(document, general)

    add_heading(document, "Candidatura y proceso electoral", 2)
    add_records(document, candidate.get("candidaturas"))
    if candidate.get("procesos"):
        add_heading(document, "Procesos asociados", 3)
        add_records(document, candidate.get("procesos"))

    sections = [
        ("Educación y títulos", detail.get("educacion")),
        ("Trayectoria política", detail.get("trayectoria")),
        ("Afiliaciones políticas", detail.get("afiliaciones")),
    ]
    for heading, content in sections:
        add_heading(document, heading, 2)
        add_records(document, content)

    contracts = detail.get("contratos") or {}
    add_heading(document, "Contratos, actividad económica y empresas", 2)
    add_heading(document, "Información tributaria SUNAT", 3)
    add_records(document, contracts.get("sunat"))
    add_heading(document, "Contratos con el Estado", 3)
    add_records(document, contracts.get("govContracts"))
    add_heading(document, "Derechos mineros", 3)
    add_records(document, contracts.get("miningRights"))
    add_heading(document, "Empresas relacionadas", 3)
    add_records(document, contracts.get("empresas"))

    add_heading(document, "Propiedades declaradas", 2)
    add_records(document, detail.get("propiedades"))

    debts = detail.get("deudas") or {}
    add_heading(document, "Deudas y obligaciones", 2)
    add_heading(document, "Deudas alimentarias", 3)
    add_records(document, debts.get("alimentarias"))
    add_heading(document, "Deudas judiciales", 3)
    add_records(document, debts.get("judiciales"))
    add_heading(document, "Licencia, infracciones y papeletas", 3)
    driving = debts.get("drivingRecord")
    if isinstance(driving, dict):
        main_driving = {key: value for key, value in driving.items() if key != "papeletas"}
        add_records(document, main_driving)
        if driving.get("papeletas"):
            add_heading(document, "Detalle de papeletas", 3)
            add_records(document, driving.get("papeletas"))
    else:
        add_records(document, driving)

    sanctions = detail.get("sanciones") or {}
    add_heading(document, "Sanciones y sentencias", 2)
    add_heading(document, "Sentencias declaradas", 3)
    add_records(document, sanctions.get("sentencias"))
    add_heading(document, "Registro Nacional de Abogados Sancionados", 3)
    add_records(document, sanctions.get("rnas"))
    add_heading(document, "Sanciones SERVIR", 3)
    add_records(document, sanctions.get("servir"))
    add_heading(document, "Sanciones OECE", 3)
    add_records(document, sanctions.get("oeceSanctions"))

    add_heading(document, "Experiencia laboral", 2)
    add_records(document, detail.get("experienciaLaboral"))

    add_heading(document, "Indicadores publicados por la fuente", 2)
    add_records(document, detail.get("indicadores"))

    add_heading(document, "Otros datos del registro", 2)
    add_key_value_table(document, {
        "proceso": record.get("proceso"),
        "es_finalista": record.get("es_finalista"),
        "companeros_formula": record.get("companeros_formula"),
    })

    if number < total:
        document.add_page_break()


def main() -> None:
    payload = json.loads(SOURCE.read_text(encoding="utf-8"))
    metadata = payload.get("metadata") or {}
    candidates = payload.get("candidatos") or []

    document = Document()
    configure_document(document)
    document.core_properties.title = "Candidatos a la Alcaldía de Lima - ERM 2026"
    document.core_properties.subject = "Perfiles de candidatos a alcalde provincial de Lima"
    document.core_properties.author = "Informe generado a partir de Revisa Tu Candidato"
    document.core_properties.keywords = "Lima, alcaldes, candidatos, ERM 2026"

    add_cover(document, metadata)
    add_summary(document, candidates)
    for number, record in enumerate(candidates, 1):
        add_candidate(document, record, number, len(candidates))

    document.save(OUTPUT)
    print(OUTPUT)
    print(f"Candidatos incluidos: {len(candidates)}")


if __name__ == "__main__":
    main()
